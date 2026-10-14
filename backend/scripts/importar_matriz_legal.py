import argparse
import os
import re
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from db.connection import get_cursor, init_pool  

HOJA_A_TIPO = {
    "1. Ley": "Ley",
    "2. Decreto": "Decreto",
    "3. Resolución": "Resolución",
    "4. Circular": "Circular",
    "5. Guía": "Guía",
}

FILA_INICIO_DATOS = 7  

def extraer_numero(texto_norma: str) -> str:
    """Extrae el número de la norma a partir del texto de la columna 'Norma'.
    Ej: 'Ley 9' -> '9'; 'Constitución Política...' -> 'S/N'."""
    coincidencia = re.search(r"(\d[\d.\-/]*)", texto_norma)
    return coincidencia.group(1) if coincidencia else "S/N"


def obtener_o_crear_id(cur, tabla: str, columna_id: str, columna_nombre: str, valor: str) -> int:
    valor = (valor or "Sin especificar").strip()[:250]
    cur.execute(f"SELECT {columna_id} FROM {tabla} WHERE {columna_nombre} = %s", (valor,))
    fila = cur.fetchone()
    if fila:
        return fila[columna_id]
    cur.execute(
        f"INSERT INTO {tabla} ({columna_nombre}) VALUES (%s) RETURNING {columna_id}",
        (valor,),
    )
    return cur.fetchone()[columna_id]


def importar_hoja(cur, hoja, nombre_tipo: str, id_usuario: int) -> int:
    cur.execute("SELECT id_tipo_norma FROM tipo_norma WHERE nombre_tipo = %s", (nombre_tipo,))
    id_tipo_norma = cur.fetchone()["id_tipo_norma"]
    cur.execute("SELECT id_estado_vigencia FROM estado_vigencia WHERE nombre_estado = 'Vigente'")
    id_vigente = cur.fetchone()["id_estado_vigencia"]

    contador = 0
    for fila in hoja.iter_rows(min_row=FILA_INICIO_DATOS, values_only=True):
        titulo = fila[0]
        if not titulo:
            continue
        anio = fila[1] or 1900
        emisor = fila[2] or "Sin especificar"
        articulos = fila[3]
        descripcion = fila[4]
        exigencia = fila[5]
        registros_evidencia = fila[6]
        frecuencia = fila[7]
        responsable_cargo = fila[8]

        id_entidad = obtener_o_crear_id(cur, "entidad_emisora", "id_entidad", "nombre", emisor)
        numero = extraer_numero(str(titulo))

        cur.execute(
            "SELECT 1 FROM norma WHERE id_tipo_norma = %s AND numero = %s AND anio = %s",
            (id_tipo_norma, numero, int(anio)),
        )
        if cur.fetchone():
            continue  
        cur.execute(
            """
            INSERT INTO norma (id_tipo_norma, numero, anio, id_entidad_emisora, titulo,
                                descripcion, articulos_que_aplican, exigencia_legal,
                                registros_evidencia_tipo, id_estado_vigencia,
                                responsable_cargo, id_usuario_responsable)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (id_tipo_norma, numero, int(anio), id_entidad, str(titulo)[:400],
             descripcion, articulos, exigencia, registros_evidencia, id_vigente,
             str(responsable_cargo)[:200] if responsable_cargo else None, id_usuario),
        )
        contador += 1
    return contador


def main():
    parser = argparse.ArgumentParser(description="Importar matriz legal SST a NormaLAB.")
    parser.add_argument("ruta_excel", help="Ruta al archivo .xlsx de la matriz legal")
    parser.add_argument("--usuario", required=True, help="nombre_usuario que quedará como responsable")
    args = parser.parse_args()

    init_pool()
    with get_cursor(commit=True) as cur:
        cur.execute("SELECT id_usuario FROM usuario WHERE nombre_usuario = %s", (args.usuario,))
        fila = cur.fetchone()
        if not fila:
            print(f"No existe el usuario '{args.usuario}'. Créelo primero (ver README).")
            sys.exit(1)
        id_usuario = fila["id_usuario"]

        libro = openpyxl.load_workbook(args.ruta_excel, data_only=True)
        total = 0
        for nombre_hoja, nombre_tipo in HOJA_A_TIPO.items():
            if nombre_hoja not in libro.sheetnames and f"{nombre_hoja} " not in libro.sheetnames:
                continue
            hoja_real = nombre_hoja if nombre_hoja in libro.sheetnames else f"{nombre_hoja} "
            hoja = libro[hoja_real]
            insertadas = importar_hoja(cur, hoja, nombre_tipo, id_usuario)
            print(f"{nombre_tipo}: {insertadas} normas importadas.")
            total += insertadas

        print(f"\nTotal importado: {total} normas.")


if __name__ == "__main__":
    main()
