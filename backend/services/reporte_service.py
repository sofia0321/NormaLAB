"""RF-009 / RF-010 (y exportación de RF-003) — Reportes en Excel y PDF.

Cada exportación: (1) valida que haya datos, (2) genera el archivo, (3) lo guarda en disco,
(4) lo registra en `reporte_generado` y (5) deja un evento EXPORTACION en el log de auditoría.
"""
import json
import os
from datetime import datetime
from io import BytesIO
from typing import Optional
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from config.settings import APP_CONFIG
from db.connection import get_cursor
from models.entities import LogAuditoria
from repositories.auditoria_repository import AuditoriaRepository
from repositories.norma_repository import NormaRepository
from services.auth_service import AuthService
from services.norma_service import ValidacionError

FORMATOS = {"xlsx": "XLSX", "pdf": "PDF"}
TIPOS_REPORTE = {"matriz": "Matriz normativa", "incumplimientos": "Normas en incumplimiento"}

# (clave en matriz_expandida, encabezado, ancho Excel, ancho relativo PDF)
COLUMNAS = [
    ("tipo", "Tipo", 12, 1.3), ("numero", "Número", 10, 1.2), ("anio", "Año", 7, 0.9),
    ("entidad_emisora", "Entidad emisora", 22, 2.2), ("descripcion", "Descripción", 45, 4.0),
    ("exigencia_legal", "Exigencia legal", 45, 3.6), ("frecuencia", "Frecuencia", 12, 1.5),
    ("responsable", "Responsable", 20, 2.0), ("cumplimiento", "Cumplimiento", 14, 1.7),
    ("fecha_evaluacion", "Fecha evaluación", 14, 1.6),
    ("acciones_realizadas", "Acciones realizadas", 40, 3.2),
]
COLUMNAS_PDF = [c for c in COLUMNAS if c[0] != "acciones_realizadas"]  # PDF más legible
COLUMNAS_XLSX = COLUMNAS + [("articulos", "Artículos aplicables", 45, 0),
                            ("acciones_abiertas", "Acciones correctivas abiertas", 30, 0)]


def _celda(valor) -> str:
    if valor is None or valor == "":
        return ""
    if hasattr(valor, "strftime"):
        return valor.strftime("%d/%m/%Y")  # formato DD/MM/AAAA del documento de RF
    return str(valor)


def filtrar_filas(filas: list[dict], f: dict) -> list[dict]:
    """Mismos filtros que la pantalla 'Matriz completa' (para que el reporte = lo que se ve)."""
    out = list(filas)
    if f.get("tipo"):
        out = [r for r in out if str(r.get("tipo", "")).lower() == f["tipo"].lower()]
    if f.get("anio"):
        out = [r for r in out if r.get("anio") == int(f["anio"])]
    if f.get("numero"):
        out = [r for r in out if f["numero"].lower() in str(r.get("numero", "")).lower()]
    if f.get("entidad"):
        out = [r for r in out if f["entidad"].lower() in str(r.get("entidad_emisora", "")).lower()]
    if f.get("cumplimiento"):
        out = [r for r in out if str(r.get("cumplimiento", "")).lower() == f["cumplimiento"].lower()]
    return out


class ReporteService:

    def __init__(self, auth_service: AuthService):
        self._auth = auth_service
        self._normas = NormaRepository()
        self._auditoria = AuditoriaRepository()

    # ---- API pública ------------------------------------------------------
    def exportar(self, formato: str, tipo_reporte: str, filtros: dict) -> tuple[str, str]:
        """Genera el reporte y devuelve (ruta_absoluta, nombre_descarga)."""
        if formato not in FORMATOS:
            raise ValidacionError("Seleccione un formato de exportación válido (Excel o PDF).")
        if tipo_reporte not in TIPOS_REPORTE:
            raise ValidacionError("Seleccione un tipo de reporte válido.")
        filtros = {k: (v or "").strip() for k, v in filtros.items()}
        if filtros.get("anio") and not filtros["anio"].isdigit():
            raise ValidacionError("El año del filtro debe ser numérico.")
        if tipo_reporte == "incumplimientos":
            filtros["cumplimiento"] = "No Cumple"

        filas = filtrar_filas(self._normas.matriz_expandida(), filtros)
        if not filas:
            raise ValidacionError("No hay registros disponibles para exportar.")

        usuario = self._auth.obtener_usuario_actual()
        ahora = datetime.now()
        titulo = TIPOS_REPORTE[tipo_reporte]
        contenido = (self._xlsx(filas, titulo, filtros, usuario.nombre_completo, ahora)
                     if formato == "xlsx"
                     else self._pdf(filas, titulo, filtros, usuario.nombre_completo, ahora))

        nombre = f"{tipo_reporte}_{ahora:%Y%m%d_%H%M%S}.{formato}"
        carpeta = self.carpeta_reportes()
        ruta = os.path.join(carpeta, nombre)
        with open(ruta, "wb") as fh:
            fh.write(contenido)

        id_reporte = self._registrar(tipo_reporte, formato, ruta, usuario.id_usuario, filtros, len(filas))
        self._auditoria.registrar(LogAuditoria(
            id_log=None, id_usuario=usuario.id_usuario, tipo_operacion="EXPORTACION",
            entidad_afectada="reporte", id_entidad_afectada=id_reporte,
            descripcion_registro=f"Exportación de «{titulo}» en {FORMATOS[formato]} ({len(filas)} registro(s))",
            detalle={"formato": FORMATOS[formato], "tipo_reporte": tipo_reporte,
                     "filtros": {k: v for k, v in filtros.items() if v}, "registros": len(filas)},
        ))
        return ruta, nombre

    def listar_generados(self, limite: int = 20) -> list[dict]:
        with get_cursor() as cur:
            cur.execute(
                """
                SELECT r.id_reporte, r.tipo_reporte, r.formato, r.fecha_generacion,
                       r.total_registros, r.filtros, u.nombre_completo
                FROM reporte_generado r JOIN usuario u ON u.id_usuario = r.id_usuario
                ORDER BY r.fecha_generacion DESC, r.id_reporte DESC LIMIT %s
                """, (limite,))
            return [dict(r) for r in cur.fetchall()]

    def ruta_de(self, id_reporte: int) -> Optional[tuple[str, str]]:
        with get_cursor() as cur:
            cur.execute("SELECT ruta_archivo FROM reporte_generado WHERE id_reporte = %s", (id_reporte,))
            row = cur.fetchone()
        if not row:
            return None
        ruta = os.path.realpath(row["ruta_archivo"])
        base = os.path.realpath(self.carpeta_reportes())
        # Defensa: solo se sirven archivos que estén dentro de la carpeta de reportes.
        if not ruta.startswith(base + os.sep) or not os.path.isfile(ruta):
            return None
        return ruta, os.path.basename(ruta)

    @staticmethod
    def carpeta_reportes() -> str:
        carpeta = os.path.join(APP_CONFIG.upload_dir, "reportes")
        os.makedirs(carpeta, exist_ok=True)
        return carpeta

    # ---- Persistencia -----------------------------------------------------
    @staticmethod
    def _registrar(tipo, formato, ruta, id_usuario, filtros, total) -> int:
        with get_cursor(commit=True) as cur:
            cur.execute(
                """
                INSERT INTO reporte_generado
                    (tipo_reporte, formato, ruta_archivo, id_usuario, filtros, total_registros)
                VALUES (%s, %s, %s, %s, %s, %s) RETURNING id_reporte
                """,
                (tipo, FORMATOS[formato], ruta, id_usuario,
                 json.dumps({k: v for k, v in filtros.items() if v}), total))
            return cur.fetchone()["id_reporte"]

    # ---- Excel (RF-009) ---------------------------------------------------
    @staticmethod
    def _xlsx(filas, titulo, filtros, usuario, ahora) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Matriz normativa"[:31]
        borde = Side(style="thin", color="C9D1DC")
        n = len(COLUMNAS_XLSX)

        ws.cell(row=1, column=1, value=f"{APP_CONFIG.institucion} — NormaLAB").font = Font(bold=True, size=13)
        ws.cell(row=2, column=1, value=titulo).font = Font(bold=True, size=11)
        activos = "; ".join(f"{k}: {v}" for k, v in filtros.items() if v) or "ninguno"
        ws.cell(row=3, column=1,
                value=f"Generado el {ahora:%d/%m/%Y %H:%M} por {usuario} · Filtros: {activos} · "
                      f"Registros: {len(filas)}").font = Font(italic=True, size=9, color="555555")

        for j, (_, enc, ancho, _) in enumerate(COLUMNAS_XLSX, start=1):
            c = ws.cell(row=5, column=j, value=enc)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E79")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = Border(top=borde, bottom=borde, left=borde, right=borde)
            ws.column_dimensions[get_column_letter(j)].width = ancho

        for i, fila in enumerate(filas, start=6):
            for j, (clave, *_r) in enumerate(COLUMNAS_XLSX, start=1):
                c = ws.cell(row=i, column=j, value=_celda(fila.get(clave)))
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.border = Border(top=borde, bottom=borde, left=borde, right=borde)

        ws.freeze_panes = "A6"
        ws.auto_filter.ref = f"A5:{get_column_letter(n)}{5 + len(filas)}"
        ws.sheet_view.zoomScale = 90
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # ---- PDF (RF-010) -----------------------------------------------------
    @staticmethod
    def _pdf(filas, titulo, filtros, usuario, ahora) -> bytes:
        buf = BytesIO()
        doc = SimpleDocTemplate(
            buf, pagesize=landscape(A4), leftMargin=1.2 * cm, rightMargin=1.2 * cm,
            topMargin=3.0 * cm, bottomMargin=1.5 * cm, title=titulo,
            author=APP_CONFIG.app_name)
        est = getSampleStyleSheet()
        celda = ParagraphStyle("c", parent=est["BodyText"], fontSize=7, leading=8.5)
        cab = ParagraphStyle("h", parent=celda, textColor=colors.white, fontName="Helvetica-Bold",
                             alignment=1)

        def p(texto, estilo=celda):
            return Paragraph(escape(_celda(texto)).replace("\n", "<br/>"), estilo)

        datos = [[p(c[1], cab) for c in COLUMNAS_PDF]]
        datos += [[p(f.get(c[0])) for c in COLUMNAS_PDF] for f in filas]
        disponible = landscape(A4)[0] - 2.4 * cm
        peso = sum(c[3] for c in COLUMNAS_PDF)
        tabla = Table(datos, colWidths=[disponible * c[3] / peso for c in COLUMNAS_PDF], repeatRows=1)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D1DC")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6FA")]),
        ]))
        activos = "; ".join(f"{k}: {v}" for k, v in filtros.items() if v) or "ninguno"
        sub = ParagraphStyle("s", parent=est["BodyText"], fontSize=8, textColor=colors.HexColor("#555555"))

        def encabezado_pie(canvas, _doc):
            canvas.saveState()
            ancho, alto = landscape(A4)
            canvas.setFont("Helvetica-Bold", 12)
            canvas.drawString(1.2 * cm, alto - 1.3 * cm, APP_CONFIG.institucion)
            canvas.setFont("Helvetica", 10)
            canvas.drawString(1.2 * cm, alto - 1.9 * cm, f"NormaLAB — {titulo}")
            canvas.setFont("Helvetica", 8)
            canvas.setFillColor(colors.HexColor("#555555"))
            canvas.drawRightString(ancho - 1.2 * cm, alto - 1.3 * cm,
                                   f"Fecha de generación: {ahora:%d/%m/%Y %H:%M}")
            canvas.drawRightString(ancho - 1.2 * cm, alto - 1.9 * cm, f"Generado por: {usuario}")
            canvas.setStrokeColor(colors.HexColor("#1F4E79"))
            canvas.line(1.2 * cm, alto - 2.3 * cm, ancho - 1.2 * cm, alto - 2.3 * cm)
            canvas.drawRightString(ancho - 1.2 * cm, 0.8 * cm, f"Página {canvas.getPageNumber()}")
            canvas.restoreState()

        doc.build([Paragraph(f"Filtros aplicados: {escape(activos)} · Registros: {len(filas)}", sub),
                   Spacer(1, 0.2 * cm), tabla],
                  onFirstPage=encabezado_pie, onLaterPages=encabezado_pie)
        return buf.getvalue()
